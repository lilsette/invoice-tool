# /// script
# requires-python = ">=3.10"
# dependencies = ["anthropic>=0.70", "openpyxl>=3.1"]
# ///
"""Turn supplier PDF invoices into a filled-in Retail Express
"Purchase Orders - Purchase Order Details" import sheet.

    uv run invoice_to_po.py INVOICE.pdf [MORE.pdf ...]

Retail Express makes a different template per supplier + outlet (Inventory → purchase
order import → Get Excel Template). Save each one into templates/ and the tool picks
the matching one per invoice (or pass --template). The template's master sheet is used
to match invoice lines to Retail Express ProductIds; lines it can't match are flagged.

Each invoice becomes its own .xlsx, written next to the PDF unless --out-dir is given.
"""

from __future__ import annotations

import argparse
import base64
import json
import sys
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

MODEL = "claude-opus-5-5"

# Template columns, in order, as Retail Express names them.
COLUMNS = [
    "ProductId",
    "ManufacturerSKU",
    "SupplierSKU",
    "ShortDescription",
    "SupplierSKU2",
    "Qty",
    "SupplierBuy",
]

EXTRACTION_PROMPT = """\
You are reading a supplier invoice so its lines can be loaded into a Retail Express
purchase order. Return every product line on the invoice.

For each line:
- supplier_sku: the supplier's own product code / item number / SKU exactly as printed
  (keep leading zeros, dashes and letters). Empty string if the line has none.
- barcode: the EAN/UPC/APN barcode if printed, else empty string.
- description: the product description as printed (trim to 50 characters).
- qty: the number of SELLABLE UNITS supplied, as a whole number. If the invoice bills in
  cartons/packs and states the units per carton, multiply it out (e.g. 2 ctn x 12 = 24)
  and say so in note.
- unit_cost_ex_tax: the cost of ONE sellable unit EXCLUDING GST/VAT, after any line
  discount, as a decimal number. If the invoice only shows tax-inclusive prices, remove
  the tax at the rate the invoice states. If qty was multiplied out from cartons, divide
  the carton price the same way.
- note: anything a buyer should double-check (unclear quantity, free goods, backorder,
  pack conversion). Empty string if nothing.

Skip lines that are not stock: freight, delivery, surcharges, deposits, rounding,
subtotals, payment terms. Do include free/bonus stock lines at cost 0 with a note.

Also return the header fields: supplier_name, invoice_number, invoice_date (YYYY-MM-DD
or as printed), and subtotal_ex_tax for the stock lines if the invoice shows one
(null if not shown).
"""

EXTRACTION_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["supplier_name", "invoice_number", "invoice_date", "subtotal_ex_tax", "lines"],
    "properties": {
        "supplier_name": {"type": "string"},
        "invoice_number": {"type": "string"},
        "invoice_date": {"type": "string"},
        "subtotal_ex_tax": {"type": ["number", "null"]},
        "lines": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["supplier_sku", "barcode", "description", "qty", "unit_cost_ex_tax", "note"],
                "properties": {
                    "supplier_sku": {"type": "string"},
                    "barcode": {"type": "string"},
                    "description": {"type": "string"},
                    "qty": {"type": "integer"},
                    "unit_cost_ex_tax": {"type": "number"},
                    "note": {"type": "string"},
                },
            },
        },
    },
}


@dataclass
class Line:
    supplier_sku: str
    barcode: str
    description: str
    qty: int
    unit_cost: Decimal
    note: str
    product_id: str = ""     # from the template's master sheet, when the line matched
    matched_on: str = ""     # which field matched, for the check sheet


@dataclass
class Invoice:
    supplier_name: str
    invoice_number: str
    invoice_date: str
    subtotal_ex_tax: Decimal | None
    lines: list[Line]


# --------------------------------------------------------------------------- extraction


def extract_invoice(pdf_bytes: bytes, name: str = "invoice") -> dict:
    """Ask Claude to read the PDF and return the schema above as a dict."""
    import anthropic

    client = anthropic.Anthropic()
    pdf_b64 = base64.standard_b64encode(pdf_bytes).decode()
    with client.beta.messages.stream(
        model=MODEL,
        max_tokens=64000,
        betas=["server-side-fallback-2026-07-01"],
        extra_body={"fallbacks": "default"},
        output_config={"effort": "medium", "format": {"type": "json_schema", "schema": EXTRACTION_SCHEMA}},
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": pdf_b64}},
                    {"type": "text", "text": EXTRACTION_PROMPT},
                ],
            }
        ],
    ) as stream:
        response = stream.get_final_message()

    if response.stop_reason == "refusal":
        raise RuntimeError(f"{name}: the model declined to read this file")
    if response.stop_reason == "max_tokens":
        raise RuntimeError(f"{name}: invoice too long for one pass — split the PDF")
    text = "".join(b.text for b in response.content if b.type == "text")
    return json.loads(text)


def to_invoice(data: dict) -> Invoice:
    def money(v) -> Decimal:
        return Decimal(str(v)).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)

    subtotal = data.get("subtotal_ex_tax")
    return Invoice(
        supplier_name=data.get("supplier_name", ""),
        invoice_number=data.get("invoice_number", ""),
        invoice_date=data.get("invoice_date", ""),
        subtotal_ex_tax=None if subtotal is None else money(subtotal),
        lines=[
            Line(
                supplier_sku=str(l.get("supplier_sku", "")).strip(),
                barcode=str(l.get("barcode", "")).strip(),
                description=str(l.get("description", "")).strip()[:50],
                qty=int(l["qty"]),
                unit_cost=money(l["unit_cost_ex_tax"]),
                note=str(l.get("note", "")).strip(),
            )
            for l in data["lines"]
        ],
    )


# --------------------------------------------------------------------------- sheet


def find_cell(ws: Worksheet, label: str):
    want = label.replace(" ", "").lower()
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.replace(" ", "").lower() == want:
                return cell
    return None


def build_template() -> Workbook:
    """A stand-in for the Retail Express template, laid out like the original."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Order Details"
    navy = PatternFill("solid", fgColor="000080")
    red = PatternFill("solid", fgColor="D9341F")
    dark = PatternFill("solid", fgColor="333333")
    white_bold = Font(bold=True, color="FFFFFF")

    ws["A1"] = "Purchase Orders - Purchase Order Details"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    for col in "ABCD":
        ws[f"{col}1"].fill = navy
    ws["A3"], ws["A4"] = "Outlet", "Supplier"
    for r in (3, 4):
        ws[f"A{r}"].font = Font(bold=True)
        ws[f"B{r}"].fill = red
        ws[f"B{r}"].font = Font(color="FFFFFF")

    types = [
        ("Integer", "Whole Number"),
        ("Reference Only", "See Master Sheet"),
        ("Text (lookup)", "50 chars"),
        ("Reference Only", "See Master Sheet"),
        ("Text (lookup)", "50 chars"),
        ("Integer", "Whole Number"),
        ("Money", "$123.45"),
    ]
    required = {"ProductId", "SupplierSKU", "SupplierSKU2", "Qty"}
    for i, (name, (t1, t2)) in enumerate(zip(COLUMNS, types), start=1):
        ws.cell(6, i, t1)
        ws.cell(7, i, t2)
        if t1 == "Reference Only":
            for r in (6, 7):
                ws.cell(r, i).fill = dark
                ws.cell(r, i).font = Font(color="FFFFFF")
        h = ws.cell(9, i, name)
        h.font = white_bold if name in required else Font(bold=True)
        if name in required:
            h.fill = red
        ws.column_dimensions[h.column_letter].width = 22
    ws["F1"] = (
        "Enter the Product ID and/or SupplierSKU and/or Supplier SKU 2 along with a desired "
        "Quantity for each product. To the left, please ensure you have entered at minimum an "
        "Outlet and a Supplier Code into the boxes provided to create your purchase order."
    )
    ws["F1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("F1:J4")
    return wb


def norm(v) -> str:
    """Compare codes loosely: case, spaces, dashes and dots don't matter."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "".join(ch for ch in str(v).upper() if ch.isalnum())


def header_key(v) -> str:
    return norm(v).lower()


def po_sheet(wb: Workbook) -> Worksheet:
    for sheet in wb.worksheets:
        if find_cell(sheet, "SupplierSKU") and find_cell(sheet, "Qty"):
            return sheet
    raise RuntimeError("template has no sheet with 'SupplierSKU' and 'Qty' headers")


def box_value(ws: Worksheet, label: str):
    cell = find_cell(ws, label)
    return None if cell is None else ws.cell(cell.row, cell.column + 1).value


# Master-sheet columns worth matching on, by normalised header.
SKU_HEADERS = ("suppliersku", "suppliersku2", "manufacturersku", "sku")
BARCODE_HEADERS = ("barcode", "ean", "apn", "upc", "gtin", "productbarcode")


class Master:
    """The product list Retail Express puts in a supplier's template ("See Master Sheet")."""

    def __init__(self, rows: list[dict[str, object]]):
        self.rows = rows
        self.index: dict[str, dict[str, dict]] = {}
        for row in rows:
            for col, value in row.items():
                key = norm(value)
                if key and (col in SKU_HEADERS or col in BARCODE_HEADERS):
                    self.index.setdefault(col, {}).setdefault(key, row)
                    self.index.setdefault(col, {}).setdefault(key.lstrip("0"), row)

    @classmethod
    def from_workbook(cls, wb: Workbook, skip: Worksheet) -> "Master | None":
        for ws in wb.worksheets:
            if ws is skip:
                continue
            for header_row in ws.iter_rows(max_row=min(ws.max_row, 30)):
                heads = {header_key(c.value): c.column for c in header_row if c.value is not None}
                if "productid" in heads and any(h in heads for h in SKU_HEADERS + BARCODE_HEADERS):
                    rows = []
                    for r in ws.iter_rows(min_row=header_row[0].row + 1, values_only=True):
                        row = {h: r[c - 1] for h, c in heads.items() if c - 1 < len(r)}
                        if row.get("productid") not in (None, ""):
                            rows.append(row)
                    return cls(rows) if rows else None
        return None

    def match(self, line: Line) -> tuple[dict, str] | None:
        tries = [(line.supplier_sku, SKU_HEADERS + BARCODE_HEADERS),
                 (line.barcode, BARCODE_HEADERS + SKU_HEADERS)]
        for value, cols in tries:
            key = norm(value)
            if not key:
                continue
            for col in cols:
                idx = self.index.get(col, {})
                hit = idx.get(key) or idx.get(key.lstrip("0"))
                if hit:
                    return hit, col
        return None


def apply_master(master: Master | None, inv: Invoice) -> None:
    if master is None:
        return
    for line in inv.lines:
        found = master.match(line)
        if found:
            row, col = found
            pid = row["productid"]
            line.product_id = str(int(pid)) if isinstance(pid, float) and pid.is_integer() else str(pid)
            line.matched_on = col
            if row.get("suppliersku"):
                line.supplier_sku = str(row["suppliersku"])


def fill(ws: Worksheet, outlet: str | None, supplier_code: str | None, inv: Invoice, sku2: str) -> None:
    # The Retail Express template arrives with Outlet and Supplier already filled in;
    # only overwrite them when told to, or when they are blank.
    for label, value in (("Outlet", outlet), ("Supplier", supplier_code)):
        cell = find_cell(ws, label)
        if cell is None:
            raise RuntimeError(f"template has no '{label}' label cell")
        box = ws.cell(cell.row, cell.column + 1)
        if value:
            box.value = value
        elif box.value in (None, ""):
            raise RuntimeError(f"the template's {label} box is empty — set the {label.lower()}")

    header = find_cell(ws, "SupplierSKU")
    cols = {header_key(c.value): c.column for c in ws[header.row] if c.value is not None}

    first = header.row + 1
    # Clear anything left in the template's data area.
    for r in range(first, ws.max_row + 1):
        for c in cols.values():
            ws.cell(r, c).value = None

    # Backordered / nil-supplied lines stay on the check sheet but not on the PO.
    for i, line in enumerate(l for l in inv.lines if l.qty > 0):
        r = first + i
        if line.product_id and "productid" in cols:
            ws.cell(r, cols["productid"], int(line.product_id) if line.product_id.isdigit() else line.product_id)
        ws.cell(r, cols["suppliersku"], line.supplier_sku or None)
        if sku2 == "barcode" and "suppliersku2" in cols:
            ws.cell(r, cols["suppliersku2"], line.barcode or None)
        if "shortdescription" in cols:
            ws.cell(r, cols["shortdescription"], line.description)
        ws.cell(r, cols["qty"], line.qty)
        if "supplierbuy" in cols:
            c = ws.cell(r, cols["supplierbuy"], float(line.unit_cost))
            c.number_format = '"$"#,##0.00##'


def add_check_sheet(wb: Workbook, inv: Invoice, source: str, has_master: bool) -> Decimal:
    ws = wb.create_sheet("Invoice Check")
    ws.append(["Source", source])
    ws.append(["Supplier on invoice", inv.supplier_name])
    ws.append(["Invoice #", inv.invoice_number])
    ws.append(["Invoice date", inv.invoice_date])
    total = lines_total(inv)
    ws.append(["Lines total ex tax", float(total)])
    ws.append(["Invoice subtotal ex tax", None if inv.subtotal_ex_tax is None else float(inv.subtotal_ex_tax)])
    ws.append([])
    ws.append(["ProductId", "SupplierSKU", "Barcode", "Description", "Qty", "Unit ex tax", "Line ex tax",
               "Matched on", "Note"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    amber = PatternFill("solid", fgColor="FFE699")
    for l in inv.lines:
        ws.append([l.product_id, l.supplier_sku, l.barcode, l.description, l.qty, float(l.unit_cost),
                   float((l.unit_cost * l.qty).quantize(Decimal("0.01"))),
                   l.matched_on or ("NOT IN RETAIL EXPRESS" if has_master else ""), l.note])
        if (has_master and not l.product_id) or l.qty <= 0:
            for c in ws[ws.max_row]:
                c.fill = amber
    for col, w in zip("ABCDEFGHI", (10, 18, 16, 40, 8, 12, 12, 22, 50)):
        ws.column_dimensions[col].width = w
    return total


def lines_total(inv: Invoice) -> Decimal:
    return sum((l.unit_cost * l.qty for l in inv.lines), Decimal(0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


# --------------------------------------------------------------------------- templates


def load_suppliers(path: Path) -> dict[str, str]:
    return json.loads(path.read_text()) if path.exists() else {}


def code_for(inv: Invoice, suppliers: dict[str, str]) -> str | None:
    name = inv.supplier_name.lower()
    for key, code in suppliers.items():
        if key.lower() in name:
            return code
    return None


def template_index(folder: Path | None) -> dict[str, dict]:
    """Supplier code → {path, outlet, supplier} for every template in the folder."""
    found: dict[str, dict] = {}
    if not folder or not folder.is_dir():
        return found
    for f in sorted(folder.glob("*.xls[xm]")):
        if f.name.startswith("~$"):
            continue
        try:
            ws = po_sheet(load_workbook(f, read_only=False))
        except Exception:
            continue
        supplier = box_value(ws, "Supplier")
        if supplier:
            found[norm(supplier)] = {"path": f, "supplier": str(supplier), "outlet": box_value(ws, "Outlet") or ""}
    return found


# --------------------------------------------------------------------------- one invoice → one PO


@dataclass
class PoResult:
    workbook: Workbook
    invoice: Invoice
    supplier_code: str
    template: Path | None
    has_master: bool
    total: Decimal
    warnings: list[str]
    filename: str


def build_po(data: dict, *, template: Path | None, outlet: str | None = None,
             supplier: str | None = None, sku2: str = "none", source: str = "") -> PoResult:
    """Fill one PO from one extracted invoice. `supplier`/`outlet` override the template's boxes
    (and are required when there is no template)."""
    inv = to_invoice(data)
    warnings: list[str] = []
    if template:
        wb = load_workbook(template)
    else:
        warnings.append("No Retail Express template for this supplier — used a stand-in layout, "
                        "so lines could not be matched to ProductIds")
        wb = build_template()
    ws = po_sheet(wb)
    master = Master.from_workbook(wb, skip=ws)
    apply_master(master, inv)
    fill(ws, outlet, supplier, inv, sku2)
    total = add_check_sheet(wb, inv, source, master is not None)
    code = str(box_value(ws, "Supplier") or "UNKNOWN")

    if inv.subtotal_ex_tax is not None and abs(inv.subtotal_ex_tax - total) > Decimal("0.05"):
        warnings.append(f"Lines add to ${total} but the invoice subtotal is ${inv.subtotal_ex_tax:.2f}")
    for l in inv.lines:
        label = f"{l.supplier_sku or '-'} {l.description}"
        if l.qty <= 0:
            warnings.append(f"Left off the PO (qty {l.qty}): {label}")
        elif master and not l.product_id:
            warnings.append(f"Not in Retail Express for this supplier: {label}")
        elif not l.supplier_sku and not l.product_id:
            warnings.append(f"No supplier SKU: {l.description}")
    filename = f"PO {code} {inv.invoice_number or Path(source).stem}.xlsx".replace("/", "-")
    return PoResult(wb, inv, code, template, master is not None, total, warnings, filename)


# --------------------------------------------------------------------------- cli


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("pdfs", nargs="+", type=Path, help="invoice PDFs (or .json extractions from --save-json)")
    p.add_argument("--template", type=Path, help="the Retail Express template downloaded for this supplier + outlet")
    p.add_argument("--templates", type=Path, default=Path(__file__).with_name("templates"),
                   help="folder of downloaded templates, one per supplier; the right one is picked per invoice")
    p.add_argument("--outlet", help='override the template\'s Outlet box, e.g. "The Rocks"')
    p.add_argument("--supplier", help="Retail Express supplier code, e.g. FG (picks/overrides the template)")
    p.add_argument("--out-dir", type=Path, help="where to write the .xlsx files (default: next to each PDF)")
    p.add_argument("--sku2", choices=["none", "barcode"], default="none",
                   help="put the invoice barcode in SupplierSKU2 (default: leave it blank)")
    p.add_argument("--save-json", action="store_true", help="also save the raw extraction as .json for review/re-use")
    p.add_argument("--suppliers", type=Path, default=Path(__file__).with_name("suppliers.json"),
                   help="JSON map of supplier-name fragment → supplier code")
    args = p.parse_args()
    templates = template_index(args.templates)
    suppliers = load_suppliers(args.suppliers)

    failures = 0
    for src in args.pdfs:
        print(f"{src.name}")
        try:
            if src.suffix.lower() == ".json":
                data = json.loads(src.read_text())
            else:
                data = extract_invoice(src.read_bytes(), src.name)
            out_dir = args.out_dir or src.parent
            out_dir.mkdir(parents=True, exist_ok=True)
            if args.save_json and src.suffix.lower() != ".json":
                (out_dir / f"{src.stem}.json").write_text(json.dumps(data, indent=2))

            template = args.template
            if not template:
                code = args.supplier or code_for(to_invoice(data), suppliers)
                if code and norm(code) in templates:
                    template = templates[norm(code)]["path"]
                elif templates:
                    # No fuzzy fallback on the supplier name: two-letter codes like "FG" turn up
                    # inside unrelated names, and the wrong template means the wrong supplier's PO.
                    raise RuntimeError(
                        f"no template for supplier '{code or data.get('supplier_name')}'. Download it from Retail "
                        f"Express (Get Excel Template → 1 Supplier, for 1 Outlet) into {args.templates}, or pass "
                        f"--supplier. Templates found for: {', '.join(templates)}")
            if template:
                print(f"  template: {template.name}")
            # Only override the template's own boxes when asked; a stand-in needs both.
            supplier = args.supplier or (None if template else code_for(to_invoice(data), suppliers))
            result = build_po(data, template=template, outlet=args.outlet, supplier=supplier,
                              sku2=args.sku2, source=src.name)
            out = out_dir / result.filename
            result.workbook.save(out)

            inv = result.invoice
            matched = sum(1 for l in inv.lines if l.product_id)
            print(f"  {inv.supplier_name} · invoice {inv.invoice_number} · {len(inv.lines)} lines"
                  + (f" ({matched} matched to Retail Express products)" if result.has_master else "")
                  + f" · ${result.total} ex tax → {out}")
            for w in result.warnings:
                print(f"  ! {w}", file=sys.stderr)
            for l in inv.lines:
                if l.note:
                    print(f"  · {l.supplier_sku or l.description}: {l.note}")
        except Exception as e:  # keep going with the other invoices
            failures += 1
            print(f"  ✗ {e}", file=sys.stderr)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
